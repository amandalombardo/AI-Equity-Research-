"""A pound reporter priced in London pence (RELX, Phase 5): GBP cash flows, no exchange rate, value per share in pence."""
import json
import shutil
import subprocess
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from desk import dcf, units, valuation_run

ROOT = Path(__file__).resolve().parents[1]
EUR_FIXTURE = ROOT / "tests" / "fixtures" / "sample_inputs_eur.yaml"


def gbp_brief():
    # the KO brief's structure, relabelled as a pound reporter with a London price
    brief = yaml.safe_load((ROOT / "briefs" / "KO.yaml").read_text())
    brief["company"].update({"ticker": "TEST", "price_currency": "GBp", "reporting_currency": "GBP"})
    brief["dcf"]["terminal_ronic"] = 0.15
    brief["dcf"]["bridge"].pop("contingent_liabilities_musd", None)
    brief["dcf"]["bridge"]["buybacks_after_balance_sheet_mgbp"] = None
    return brief


def gbp_inputs(tmp_path):
    # the euro fixture's illustrative values, renamed to pounds and without an exchange rate
    inp = yaml.safe_load(EUR_FIXTURE.read_text())
    inp.pop("fx_eur_per_gbp")
    inp.pop("priced_at_valuation_date_meur")
    for k in [k for k in inp if k.endswith("_meur")]:
        inp[k.replace("_meur", "_mgbp")] = inp.pop(k)
    path = tmp_path / "inputs.yaml"
    path.write_text(yaml.safe_dump(inp))
    return path


def test_gbp_is_a_reporting_currency():
    assert units.reporting_currency({"reporting_currency": "GBP"}) == "GBP"
    assert units.money_key("dividends_ex_before_valuation", "GBP") == "dividends_ex_before_valuation_mgbp"


def test_pound_run_needs_no_exchange_rate():
    req = valuation_run.extra_required(gbp_brief())
    assert "share_price_gbp_pence" in req and "dividends_ex_before_valuation_mgbp" in req
    assert not any(k.startswith("fx_") for k in req)  # pounds to pence is x100, no FX input to source


def test_pounds_to_pence():
    b = dcf.Bridge(net_debt=100)
    eq = dcf.equity_bridge(1000, b, shares_m=10, fx_per_gbp=1.0, reporting_currency="GBP")
    # equity GBP 900 / 10 shares = GBP 90 a share = 9,000p
    assert eq["per_share_gbp_pence"] == pytest.approx(9000)


def test_full_pound_run_rates_in_pence(tmp_path):
    brief_path = tmp_path / "TEST.yaml"
    brief_path.write_text(yaml.safe_dump(gbp_brief()))
    valuation_run.main(str(brief_path), str(gbp_inputs(tmp_path)), str(tmp_path / "val"))
    result = json.loads((tmp_path / "val" / "valuation.json").read_text())
    assert result["reporting_currency"] == "GBP" and result["price_currency"] == "GBp"
    eq = result["base"]["equity"]
    assert eq["per_share_gbp_pence"] == pytest.approx(eq["per_share_gbp"] * 100)
    assert result["base"]["rating"]["fair_value"] == pytest.approx(eq["per_share_gbp_pence"])
    ws = load_workbook(tmp_path / "val" / "dcf.xlsx")["Inputs"]
    fx_row = next(r for r in ws.iter_rows(values_only=True) if r[0] == "GBP per GBP")
    assert fx_row[1] == 1.0 and "no exchange rate" in fx_row[2]


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")
def test_pound_excel_formulas_match_python(tmp_path):
    brief_path = tmp_path / "TEST.yaml"
    brief_path.write_text(yaml.safe_dump(gbp_brief()))
    valuation_run.main(str(brief_path), str(gbp_inputs(tmp_path)), str(tmp_path))
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(tmp_path / "recalc"),
                    str(tmp_path / "dcf.xlsx")], check=True, capture_output=True, timeout=120)
    ws = load_workbook(tmp_path / "recalc" / "dcf.xlsx", data_only=True)["DCF"]
    rows = list(ws.iter_rows(values_only=True))
    start = next(i for i, r in enumerate(rows) if r[0] == "Valuation")
    checked = 0
    for r in rows[start + 1:]:
        if r[0]:
            assert r[1] == pytest.approx(r[2], rel=1e-9), r[0]
            checked += 1
    assert checked >= 13


GBP_XBRL = b"""<?xml version="1.0"?>
<xbrli:xbrl xmlns:xbrli="http://www.xbrl.org/2003/instance" xmlns:iso4217="http://www.xbrl.org/2003/iso4217"
  xmlns:ifrs-full="https://xbrl.ifrs.org/taxonomy/2025-03-27/ifrs-full" xmlns:relx="http://www.relx.com/20251231">
  <xbrli:context id="FY25"><xbrli:entity><xbrli:identifier scheme="http://www.sec.gov/CIK">0000929869</xbrli:identifier></xbrli:entity>
    <xbrli:period><xbrli:startDate>2025-01-01</xbrli:startDate><xbrli:endDate>2025-12-31</xbrli:endDate></xbrli:period></xbrli:context>
  <xbrli:context id="D25"><xbrli:entity><xbrli:identifier scheme="http://www.sec.gov/CIK">0000929869</xbrli:identifier></xbrli:entity>
    <xbrli:period><xbrli:instant>2025-12-31</xbrli:instant></xbrli:period></xbrli:context>
  <xbrli:unit id="Unit_Standard_GBP_z63PyK0v0a74"><xbrli:measure>iso4217:GBP</xbrli:measure></xbrli:unit>
  <ifrs-full:RevenueFromContractsWithCustomers contextRef="FY25" unitRef="Unit_Standard_GBP_z63PyK0v0a74" decimals="-6">9590000000</ifrs-full:RevenueFromContractsWithCustomers>
  <relx:InventoriesAndPrepublicationCosts contextRef="D25" unitRef="Unit_Standard_GBP_z63PyK0v0a74" decimals="-6">311000000</relx:InventoriesAndPrepublicationCosts>
</xbrli:xbrl>"""


def test_relx_units_and_labels_are_read_in_pounds():
    # RELX names its units "Unit_Standard_GBP_<hash>": the currency must come from the unit's measure, not its name
    from desk import data_sources
    facts = data_sources.parse_instance(GBP_XBRL)
    assert data_sources._pick(facts, "ifrs-full:RevenueFromContractsWithCustomers", "2025-12-31", unit="gbp") == 9590e6
    assert data_sources._pick(facts, "relx:InventoriesAndPrepublicationCosts", "2025-12-31", unit="gbp") == 311e6
    assert data_sources._pick(facts, "ifrs-full:RevenueFromContractsWithCustomers", "2025-12-31") is None  # not dollars
    tags, _ = data_sources.tag_map({"accounting": "ifrs"})
    assert "ifrs-full:RevenueFromContractsWithCustomers" in tags["net_sales"]
    assert "relx:InventoriesAndPrepublicationCosts" in tags["inventories"]
