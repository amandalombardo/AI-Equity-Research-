"""A euro reporter priced in London pence (Unilever, Phase 5): EUR cash flows, EUR-per-GBP conversion to pence."""
import json
import shutil
import subprocess
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from desk import dcf, units, valuation_run

ROOT = Path(__file__).resolve().parents[1]
EUR_FIXTURE = ROOT / "tests" / "fixtures" / "sample_inputs_eur.yaml"


def eur_brief():
    # the KO brief's structure, relabelled as a euro reporter with a London price and no contingent liability
    brief = yaml.safe_load((ROOT / "briefs" / "KO.yaml").read_text())
    brief["company"].update({"ticker": "TEST", "price_currency": "GBp", "reporting_currency": "EUR"})
    brief["dcf"]["terminal_ronic"] = 0.18
    brief["dcf"]["bridge"].pop("contingent_liabilities_musd", None)
    return brief


def test_reporting_currency_from_brief():
    assert units.reporting_currency({}) == "USD"  # Diageo and KO report in US dollars
    assert units.reporting_currency({"reporting_currency": "EUR"}) == "EUR"
    with pytest.raises(ValueError):
        units.reporting_currency({"reporting_currency": "JPY"})
    assert units.fx_key("EUR") == "fx_eur_per_gbp"
    assert units.per_share_key("GBp", "EUR") == "per_share_gbp_pence"
    with pytest.raises(ValueError):
        units.per_share_key("USD", "EUR")  # a dollar price needs dollar cash flows: no silent cross-rate


def test_bridge_converts_euros_to_pence():
    b = dcf.Bridge(net_debt=100, associates=50)
    eq = dcf.equity_bridge(1000, b, shares_m=10, fx_per_gbp=1.25, reporting_currency="EUR")
    # equity EUR 950 / 10 shares = EUR 95 a share; at EUR 1.25 per GBP that is GBP 76 = 7,600p
    assert eq["per_share_eur"] == pytest.approx(95)
    assert "per_share_usd" not in eq
    assert eq["per_share_gbp_pence"] == pytest.approx(7600)


def test_euro_run_needs_euro_fx_and_euro_dividends():
    req = valuation_run.extra_required(eur_brief())
    assert "fx_eur_per_gbp" in req and "share_price_gbp_pence" in req and "dividends_ex_before_valuation_meur" in req
    assert "fx_usd_per_gbp" not in req and "dividends_ex_before_valuation_musd" not in req


def test_full_euro_run_rates_in_pence(tmp_path):
    brief_path = tmp_path / "TEST.yaml"
    brief_path.write_text(yaml.safe_dump(eur_brief()))
    valuation_run.main(str(brief_path), str(EUR_FIXTURE), str(tmp_path / "val"))
    result = json.loads((tmp_path / "val" / "valuation.json").read_text())
    assert result["reporting_currency"] == "EUR" and result["price_currency"] == "GBp"
    assert "bridge_meur" in result and "enterprise_value_meur" in result["base"]
    eq = result["base"]["equity"]
    assert eq["per_share_gbp_pence"] == pytest.approx(eq["per_share_eur"] / 1.1834 * 100)
    assert result["base"]["rating"]["fair_value"] == pytest.approx(eq["per_share_gbp_pence"])
    ws = load_workbook(tmp_path / "val" / "dcf.xlsx")["DCF"]
    assert ws["A1"].value == "EUR m"


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")
def test_euro_excel_formulas_match_python(tmp_path):
    brief_path = tmp_path / "TEST.yaml"
    brief_path.write_text(yaml.safe_dump(eur_brief()))
    valuation_run.main(str(brief_path), str(EUR_FIXTURE), str(tmp_path))
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(tmp_path / "recalc"),
                    str(tmp_path / "dcf.xlsx")], check=True, capture_output=True, timeout=120)
    ws = load_workbook(tmp_path / "recalc" / "dcf.xlsx", data_only=True)["DCF"]
    rows = list(ws.iter_rows(values_only=True))
    start = next(i for i, r in enumerate(rows) if r[0] == "Valuation")
    checked = 0
    for r in rows[start + 1:]:
        if r[0]:
            assert r[1] == pytest.approx(r[2], rel=1e-9), r[0]
            checked += 1
    assert checked == 14  # includes the pence line


EUR_XBRL = b"""<?xml version="1.0"?>
<xbrli:xbrl xmlns:xbrli="http://www.xbrl.org/2003/instance" xmlns:ifrs-full="https://xbrl.ifrs.org/taxonomy/2025-03-27/ifrs-full">
  <xbrli:context id="FY25"><xbrli:entity><xbrli:identifier scheme="http://www.sec.gov/CIK">0000217410</xbrli:identifier></xbrli:entity>
    <xbrli:period><xbrli:startDate>2025-01-01</xbrli:startDate><xbrli:endDate>2025-12-31</xbrli:endDate></xbrli:period></xbrli:context>
  <xbrli:unit id="EUR"><xbrli:measure>iso4217:EUR</xbrli:measure></xbrli:unit>
  <ifrs-full:Revenue contextRef="FY25" unitRef="EUR" decimals="-6">50503000000</ifrs-full:Revenue>
</xbrli:xbrl>"""


def test_ifrs_labels_read_in_euros():
    from desk import data_sources
    tags, share_tags = data_sources.tag_map({"accounting": "ifrs"})
    assert tags["net_sales"] == ["ifrs-full:Revenue"]  # Unilever's turnover; no excise, unlike Diageo
    assert "restructuring" in tags                      # recurring at Unilever, so the Data agent reports it
    facts = data_sources.parse_instance(EUR_XBRL)
    assert data_sources._pick(facts, "ifrs-full:Revenue", "2025-12-31", unit="eur") == 50503e6
    assert data_sources._pick(facts, "ifrs-full:Revenue", "2025-12-31") is None  # never read euros as dollars
    assert units.money_key("value", "EUR") == "value_meur"  # financials_raw.csv column for a euro reporter
