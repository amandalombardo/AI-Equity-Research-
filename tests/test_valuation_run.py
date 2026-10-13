import json
import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

from desk import valuation_run

ROOT = Path(__file__).resolve().parents[1]
BRIEF = ROOT / "briefs" / "DGE.yaml"
FIXTURE = ROOT / "tests" / "fixtures" / "sample_inputs.yaml"


def test_run_refuses_missing_or_unsourced_inputs(tmp_path):
    bad = tmp_path / "inputs.yaml"
    bad.write_text("tax_rate: {value: 0.23}\n")
    with pytest.raises(SystemExit):
        valuation_run.main(str(BRIEF), str(bad), str(tmp_path / "out"))


def test_full_run_writes_outputs(tmp_path):
    valuation_run.main(str(BRIEF), str(FIXTURE), str(tmp_path))
    result = json.loads((tmp_path / "valuation.json").read_text())
    assert result["base"]["rating"]["rating"] in {"BUY", "HOLD", "SELL"}
    # bear case (no growth, lower flow-through, higher WACC) must be worth less than the base case
    assert result["bear"]["equity"]["per_share_gbp_pence"] < result["base"]["equity"]["per_share_gbp_pence"]
    assert (tmp_path / "dcf.xlsx").exists()


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")
def test_excel_formulas_match_python(tmp_path):
    valuation_run.main(str(BRIEF), str(FIXTURE), str(tmp_path))
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(tmp_path / "recalc"),
                    str(tmp_path / "dcf.xlsx")], check=True, capture_output=True, timeout=120)
    ws = load_workbook(tmp_path / "recalc" / "dcf.xlsx", data_only=True)["DCF"]
    rows = list(ws.iter_rows(values_only=True))
    start = next(i for i, r in enumerate(rows) if r[0] == "Valuation")
    checked = 0
    for r in rows[start + 1:]:  # valuation section: column B = Excel formula, column C = Python value
        if r[0]:
            assert r[1] == pytest.approx(r[2], rel=1e-9), r[0]
            checked += 1
    assert checked == 14
    # forecast block: final-year Excel formulas vs the Python check column
    for r in rows[1:start]:
        if r[0] and isinstance(r[-1], (int, float)):
            assert r[-3] == pytest.approx(r[-1], rel=1e-9), r[0]
            checked += 1
    assert checked == 18


def test_credit_spread_rule_sets_cost_of_debt():
    brief = {"dcf": {"wacc": {"credit_spread": 0.01}}}
    v = {"risk_free": 0.05, "pre_tax_cost_of_debt": 0.035}
    sources = {"pre_tax_cost_of_debt": "historic"}
    valuation_run.apply_brief_rules(brief, v, sources)
    assert v["pre_tax_cost_of_debt"] == pytest.approx(0.06)
    assert v["pre_tax_cost_of_debt_historic"] == pytest.approx(0.035)


def test_no_credit_spread_keeps_agent_input():
    v = {"risk_free": 0.05, "pre_tax_cost_of_debt": 0.035}
    valuation_run.apply_brief_rules({"dcf": {"wacc": {}}}, v, {})
    assert v["pre_tax_cost_of_debt"] == pytest.approx(0.035)


def test_balance_sheet_basis_values_full_first_year_then_rolls_forward():
    # Under valuation_basis: balance_sheet_date the DCF uses no stub (all of FY27 counted, discounted from 30 June)
    # and the equity value is grown at the cost of equity for the elapsed fraction of the year
    import yaml
    brief = yaml.safe_load(BRIEF.read_text())
    brief["dcf"]["valuation_basis"] = "balance_sheet_date"
    v, _ = valuation_run.load_inputs(str(FIXTURE), valuation_run.ROLL_FORWARD)
    out = valuation_run.build(brief, v)
    assert out["inputs"].stub_years == 0
    eq = out["equity"]
    ke = out["wacc"]["cost_of_equity"]
    assert eq["equity_value"] == pytest.approx(eq["equity_value_balance_sheet_date"] * (1 + ke) ** v["stub_years"])


def test_proportionate_minorities_scale_with_ev():
    # brief: minorities_method proportionate -> minorities = share of operating profit x EV, in base and bear alike
    import yaml
    brief = yaml.safe_load(BRIEF.read_text())
    brief["dcf"]["bridge"]["minorities_method"] = "proportionate"
    v, _ = valuation_run.load_inputs(str(FIXTURE), valuation_run.extra_required(brief))
    for out in (valuation_run.build(brief, v), valuation_run.build(brief, v, wacc_shift=0.01)):
        ev = out["dcf"]["enterprise_value"]
        assert out["equity"]["minorities_deducted"] == pytest.approx(0.03 * ev)
    brief["dcf"]["bridge"]["minorities_method"] = "book"
    assert valuation_run.build(brief, v)["equity"]["minorities_deducted"] == pytest.approx(2000)
    brief["dcf"]["bridge"]["minorities_method"] = "market"
    with pytest.raises(SystemExit):
        valuation_run.build(brief, v)
