"""A US-listed company (Coca-Cola, Phase 5): US GAAP labels, a share price in US dollars, a contingent liability."""
import json
import shutil
import subprocess
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from desk import data_sources, dcf, units, valuation_run

ROOT = Path(__file__).resolve().parents[1]
KO_BRIEF = ROOT / "briefs" / "KO.yaml"
USD_FIXTURE = ROOT / "tests" / "fixtures" / "sample_inputs_usd.yaml"

US_XBRL = b"""<?xml version="1.0"?>
<xbrl xmlns="http://www.xbrl.org/2003/instance" xmlns:us-gaap="http://fasb.org/us-gaap/2025" xmlns:ko="http://www.coca-colacompany.com/20251231">
  <context id="FY25"><entity><identifier scheme="x">1</identifier></entity>
    <period><startDate>2025-01-01</startDate><endDate>2025-12-31</endDate></period></context>
  <context id="Q4"><entity><identifier scheme="x">1</identifier></entity>
    <period><startDate>2025-10-01</startDate><endDate>2025-12-31</endDate></period></context>
  <us-gaap:Revenues contextRef="FY25" unitRef="usd" decimals="-6">47941000000</us-gaap:Revenues>
  <us-gaap:Revenues contextRef="Q4" unitRef="usd" decimals="-6">11800000000</us-gaap:Revenues>
  <ko:IncrementalTaxAndInterestLiability contextRef="FY25" unitRef="usd" decimals="-9">14000000000</ko:IncrementalTaxAndInterestLiability>
</xbrl>"""


def ko_brief(ronic=0.15):
    brief = yaml.safe_load(KO_BRIEF.read_text())
    brief["dcf"]["terminal_ronic"] = ronic  # the real brief says "pending" until the PM sets it
    return brief


def test_brief_chooses_the_us_gaap_labels():
    tags, share_tags = data_sources.tag_map({"accounting": "us_gaap"})
    assert tags["net_sales"][0] == "us-gaap:Revenues"
    assert "us-gaap:TreasuryStockCommonShares" in share_tags["treasury_shares_m"]
    assert data_sources.tag_map({})[0] is data_sources.TAGS  # Diageo's labels stay the default
    with pytest.raises(SystemExit):
        data_sources.tag_map({"accounting": "japanese_gaap"})


def test_pick_reads_full_year_us_gaap_and_company_tags():
    facts = data_sources.parse_instance(US_XBRL)
    assert data_sources._pick(facts, "us-gaap:Revenues", "2025-12-31") == 47941e6  # not the quarter
    assert data_sources._pick(facts, "ko:IncrementalTaxAndInterestLiability", "2025-12-31") == 14000e6


def test_price_currency_from_brief():
    assert units.price_currency({}) == "GBp"
    assert units.price_currency({"price_currency": "USD"}) == "USD"
    with pytest.raises(ValueError):
        units.price_currency({"price_currency": "GBP"})  # pounds, not pence: refused, it's the classic 100x trap
    assert units.per_share_key("USD") == "per_share_usd"


def test_bridge_without_fx_gives_usd_only_and_deducts_contingent_liabilities_by_hand():
    b = dcf.Bridge(net_debt=100, associates=50, contingent_liabilities=30)
    eq = dcf.equity_bridge(1000, b, shares_m=10, fx_usd_per_gbp=None)
    # equity = EV - net debt - contingent liabilities + associates = 1000 - 100 - 30 + 50
    assert eq["equity_value"] == pytest.approx(920)
    assert eq["per_share_usd"] == pytest.approx(92)
    assert eq["per_share_gbp_pence"] is None


def test_us_run_rates_in_dollars_and_deducts_the_tax_case(tmp_path):
    brief = ko_brief()
    v, _ = valuation_run.load_inputs(str(USD_FIXTURE), valuation_run.extra_required(brief))
    out = valuation_run.build(brief, v)
    assert out["rating"]["share_price"] == 90
    assert out["rating"]["fair_value"] == pytest.approx(out["equity"]["per_share_usd"])
    v_win = {**v, "contingent_liabilities": 0}
    win = valuation_run.build(brief, v_win)
    # losing the case costs exactly the liability, rolled forward like the rest of equity
    gap = win["equity"]["equity_value_balance_sheet_date"] - out["equity"]["equity_value_balance_sheet_date"]
    assert gap == pytest.approx(14900)


def test_us_run_needs_a_dollar_price_and_the_contingent_liability(tmp_path):
    brief = ko_brief()
    req = valuation_run.extra_required(brief)
    assert "share_price_usd" in req and "contingent_liabilities" in req
    assert "fx_usd_per_gbp" not in req and "share_price_gbp_pence" not in req
    bad = tmp_path / "inputs.yaml"
    bad.write_text("\n".join(l for l in USD_FIXTURE.read_text().splitlines() if not l.startswith("contingent_liabilities")))
    with pytest.raises(SystemExit):
        valuation_run.load_inputs(str(bad), req)


def test_pending_ronic_refuses_to_value(tmp_path):
    brief = ko_brief(ronic="pending")
    v, _ = valuation_run.load_inputs(str(USD_FIXTURE), valuation_run.extra_required(brief))
    with pytest.raises(SystemExit):
        valuation_run.build(brief, v)


def test_full_us_run_writes_outputs(tmp_path):
    brief_path = tmp_path / "KO.yaml"
    brief_path.write_text(yaml.safe_dump(ko_brief()))
    valuation_run.main(str(brief_path), str(USD_FIXTURE), str(tmp_path / "val"))
    result = json.loads((tmp_path / "val" / "valuation.json").read_text())
    assert result["price_currency"] == "USD"
    assert result["bear"]["equity"]["per_share_usd"] < result["base"]["equity"]["per_share_usd"]
    assert "sensitivity_value_per_share_usd" in result
    assert (tmp_path / "val" / "dcf.xlsx").exists()


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")
def test_us_excel_formulas_match_python(tmp_path):
    brief_path = tmp_path / "KO.yaml"
    brief_path.write_text(yaml.safe_dump(ko_brief()))
    valuation_run.main(str(brief_path), str(USD_FIXTURE), str(tmp_path))
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(tmp_path / "recalc"),
                    str(tmp_path / "dcf.xlsx")], check=True, capture_output=True, timeout=120)
    ws = load_workbook(tmp_path / "recalc" / "dcf.xlsx", data_only=True)["DCF"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][1] == "FY25 base"
    start = next(i for i, r in enumerate(rows) if r[0] == "Valuation")
    checked = 0
    for r in rows[start + 1:]:
        if r[0]:
            assert r[1] == pytest.approx(r[2], rel=1e-9), r[0]
            checked += 1
    assert checked == 13  # no pence line for a US-dollar price
