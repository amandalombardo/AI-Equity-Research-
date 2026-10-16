"""Risk agent, Unilever (ULVR), run 2026-10-14: the valuation beta (PM rule) and its evidence.

THE ONLY SCRIPT THAT WRITES `beta` TO inputs.yaml (Codex audit RISK-01). It follows the brief's rule
(briefs/ULVR.yaml dcf.wacc.beta): bottom-up beta = Damodaran GLOBAL Household Products unlevered beta corrected
for cash, relevered at Unilever's market D/E and the brief's 26.0% tax rate. Regressions and peer betas
(risk.json, from run_risk.py) are cross-checks only and never overwrite it.

Run from the repo root (after run_risk.py):
  PYTHONPATH=. python runs/ULVR/2026-10-14/risk/beta_evidence.py <path to betaGlobal.xlsx>
The .xlsx is the PM's betaGlobal.xls converted with: soffice --headless --convert-to xlsx --outdir <scratch> <file>
"""
import json
import re
import sys
from pathlib import Path

import openpyxl
import yaml

from desk import dcf

HERE = Path(__file__).resolve().parent
RUN = HERE.parent
REPO = RUN.parents[2]
DAMODARAN_XLS = "docs/data/betaGlobal.xls"  # uploaded by the PM 14 Oct 2026; same file as https://pages.stern.nyu.edu/~adamodar/pc/datasets/betaGlobal.xls
SHA256_XLS = "10c33f46df2a35fa7e4e8c4386c01f9da636fc2ab3e1e98307a52a947c9ab52b"
PM_VALUE = 0.774144424006051  # brief / PM message: Household Products, unlevered beta corrected for cash

brief = yaml.safe_load((REPO / "briefs" / "ULVR.yaml").read_text())
inp = yaml.safe_load((RUN / "inputs.yaml").read_text())
tax = brief["dcf"]["tax_rate"]                       # 26.0%, PM decision 9
erp = brief["dcf"]["wacc"]["equity_risk_premium"]    # 5.0%
rf = inp["risk_free"]["value"]                       # 3.18% Bund (Data agent)

# ---- 1. read the industry beta from the PM's file (do not trust a typed-in number)
wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
ws = wb["Industry Averages"]
file_date = ws["B1"].value
header = row = None
for r in ws.iter_rows(values_only=True):
    if r[0] == "Industry Name":
        header = [str(h).strip() if h else h for h in r]
    if r[0] == "Household Products":
        row = dict(zip(header, r))
if row is None:
    sys.exit("MISSING: Household Products row not found in 'Industry Averages'. STOP.")
bu = float(row["Unlevered beta corrected for cash"])
if abs(bu - PM_VALUE) > 1e-12 or row["Number of firms"] != 588 or str(file_date.date()) != "2026-01-05":
    sys.exit(f"STOP: file value {bu}, firms {row['Number of firms']}, date {file_date} do not match the PM's rule")
industry = {"industry": "Household Products", "region": "Global", "number_of_firms": row["Number of firms"],
            "file_date": str(file_date.date()), "levered_beta": row["Beta"], "industry_d_to_e": row["D/E Ratio"],
            "unlevered_beta": row["Unlevered beta"], "cash_to_firm_value": row["Cash/Firm value"],
            "unlevered_beta_corrected_for_cash": bu,
            "over_time": {h: row[h] for h in header if h and re.fullmatch(r"20\d\d|Average.*", h)},
            "source": f"Damodaran, Betas by Sector (Global), betaGlobal.xls dated 2026-01-05, sheet 'Industry Averages', "
                      f"row 'Household Products' (PM upload 14 Oct 2026, {DAMODARAN_XLS}, sha256 {SHA256_XLS[:12]}...)"}

# ---- 2. Unilever's market debt-to-equity, all in EUR m
shares_m = inp["shares_outstanding_m"]["value"]          # 2,154.33m voting shares, TVR 6-K 31 Aug 2026
price_p = inp["share_price_gbp_pence"]["value"]          # 4,710.5p, ULVR.L close 13 Oct 2026
fx = inp["fx_eur_per_gbp"]["value"]                      # 1.18190 EUR per GBP, FRED 9 Oct 2026
mcap_gbp_m = shares_m * price_p / 100                    # market cap in GBP m = shares x price in pounds
mcap_eur_m = mcap_gbp_m * fx                             # GBP -> EUR: multiply by EUR per GBP


def relever(unlevered: float, debt: float, equity: float, t: float) -> float:
    """Hamada: levered beta = unlevered beta x (1 + (1 - tax) x D/E). No desk function exists for this (FLAG)."""
    return unlevered * (1 + (1 - t) * debt / equity)


def ke(beta: float) -> float:
    """Cost of equity via desk.dcf.cost_of_equity (CAPM), for showing what each beta means."""
    return dcf.cost_of_equity(rf, beta, erp)


nd_dec = inp["net_debt"]["value"]   # 23,076, 31 Dec 2025 (FY25 results 6-K / 20-F note 15C), leases inside
nd_jun = 25961.0                    # 30 Jun 2026, H1 2026 6-K 0000217410-26-000057 ("EUR 26.0 billion"), after the 1.5bn buyback
base = relever(bu, nd_dec, mcap_eur_m, tax)
cases = {
    "BASE: cash-corrected 0.7741, net debt 31 Dec 2025": base,
    "Sensitivity: net debt 30 Jun 2026 (25,961)": relever(bu, nd_jun, mcap_eur_m, tax),
    "Sensitivity: unlevered beta without cash correction (0.7432)": relever(row["Unlevered beta"], nd_dec, mcap_eur_m, tax),
    "Sensitivity: Damodaran 2020-24 average cash-corrected (0.8752)": relever(row["Average (2020-24)"], nd_dec, mcap_eur_m, tax),
}
rj = json.loads((HERE / "risk.json").read_text())
own = rj["own_regressions_vs_ftse100"]
xchecks = {f"ULVR vs FTSE 100, {k.replace('_', ' ')}": {kk: own[k][kk] for kk in ("beta", "std_error", "ci95", "r_squared", "observations")}
           for k in own}
xchecks["Peer median, 5y monthly vs home index (raw, not relevered)"] = {"beta": rj["peer_median_beta_5y_monthly"]}
inside = {k: v["ci95"][0] <= base <= v["ci95"][1] for k, v in xchecks.items() if "ci95" in v}

out = {
    "rule": "PM, briefs/ULVR.yaml dcf.wacc.beta: bottom-up Damodaran GLOBAL Household Products unlevered beta corrected for cash, "
            "relevered at Unilever's market D/E and the 26.0% tax rate. Regressions are the cross-check.",
    "industry": industry,
    "market_cap": {"shares_m": shares_m, "price_gbp_pence": price_p, "market_cap_gbp_m": mcap_gbp_m, "fx_eur_per_gbp": fx,
                   "market_cap_eur_m": mcap_eur_m,
                   "sources": "inputs.yaml shares_outstanding_m (TVR 6-K 0001654954-26-008017), share_price_gbp_pence (Yahoo ULVR.L "
                              "close 2026-10-13), fx_eur_per_gbp (FRED DEXUSUK/DEXUSEU 2026-10-09)"},
    "net_debt_meur": {"31_dec_2025": nd_dec, "30_jun_2026": nd_jun},
    "d_to_e": {"31_dec_2025": nd_dec / mcap_eur_m, "30_jun_2026": nd_jun / mcap_eur_m},
    "tax_rate": tax,
    "relevered_betas": {k: {"beta": v, "cost_of_equity": ke(v)} for k, v in cases.items()},
    "cross_checks": xchecks, "base_inside_each_95pct_interval": inside,
    "recommended_beta": round(base, 4),
}
(HERE / "beta_evidence.json").write_text(json.dumps(out, indent=2, default=str))

# ---- 3. write ONLY `beta` into inputs.yaml, leaving every other entry's text untouched
src = (f"Risk agent, bottom-up: Damodaran GLOBAL Household Products unlevered beta corrected for cash {bu:.4f} (588 firms, "
       f"betaGlobal.xls dated 2026-01-05, sheet 'Industry Averages', PM upload 14 Oct 2026) relevered (Hamada) at Unilever's market "
       f"D/E {nd_dec / mcap_eur_m:.4f} = net debt EUR {nd_dec:,.0f}m (31 Dec 2025, FY25 results 6-K 0001654954-26-001155) / market cap "
       f"EUR {mcap_eur_m:,.1f}m ({shares_m:,.3f}m shares x {price_p:,.1f}p / 100 x {fx:.5f} EUR per GBP) and tax {tax:.1%} (brief). "
       f"Cross-checks in risk/beta_evidence.json and risk/risk.json (not used).")
text = (RUN / "inputs.yaml").read_text()
text = re.sub(r"(?ms)^beta:\n(?:[ ].*\n)*", "", text)          # drop any earlier beta block (idempotent re-runs)
text = text.rstrip("\n") + "\n" + yaml.safe_dump({"beta": {"value": float(round(base, 4)), "source": src}},
                                                   sort_keys=False, width=160)
check = yaml.safe_load(text)
assert {k: v for k, v in check.items() if k != "beta"} == {k: v for k, v in inp.items() if k != "beta"}, "other entries changed"
(RUN / "inputs.yaml").write_text(text)

print(f"Household Products bu (cash-corrected) {bu}  firms {row['Number of firms']}  date {file_date.date()}")
print(f"mcap GBP m {mcap_gbp_m:,.1f}  EUR m {mcap_eur_m:,.1f}  D/E dec {nd_dec / mcap_eur_m:.4f}  jun {nd_jun / mcap_eur_m:.4f}")
for k, v in cases.items():
    print(f"{k:65s} beta {v:.4f}  Ke {ke(v):.4%}")
print("inside 95% CI:", inside)
