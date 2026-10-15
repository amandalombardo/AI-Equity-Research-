"""Valuation agent sanity checks, rating flip points and value per 1 point of each input for RELX (run 2026-10-14).

Never edits briefs/RELX.yaml or inputs.yaml. Every variant is an in-memory copy of the engine's inputs, valued with
desk.valuation_run.build (the same code as the base case). Money in GBP m; value per share in pence (GBp) unless labelled.
Also recalculates dcf.xlsx in LibreOffice and compares every formula cell with the Python value.
Run from the repo root (after the engine and run_comps.py): PYTHONPATH=. python runs/RELX/2026-10-14/valuation/run_sanity.py
"""
import copy
import json
import shutil
import subprocess
import tempfile
from pathlib import Path

import yaml
from openpyxl import load_workbook

from desk import dcf, units
from desk.valuation_run import apply_brief_rules, build, extra_required, load_inputs

HERE = Path(__file__).resolve().parent
RUN = HERE.parent
ROOT = RUN.parents[2]
brief0 = yaml.safe_load((ROOT / "briefs" / "RELX.yaml").read_text())
v0, src = load_inputs(str(RUN / "inputs.yaml"), extra_required(brief0))
apply_brief_rules(brief0, v0, src)
comps_out = json.loads((HERE / "comps.json").read_text())
engine_json = json.loads((HERE / "valuation.json").read_text())

price = v0["share_price_gbp_pence"]
shares = v0["shares_outstanding_m"]
buy, sell = brief0["rating"]["buy_if_upside_at_least"], brief0["rating"]["sell_if_downside_at_least"]
hold_floor, buy_floor = price * (1 - sell), price * (1 + buy)   # value per share (pence) where SELL->HOLD and HOLD->BUY
ERP, G0 = brief0["dcf"]["wacc"]["equity_risk_premium"], brief0["dcf"]["terminal_growth"]
SPREAD = brief0["dcf"]["wacc"]["credit_spread"]


def run(v=None, brief=None, wacc_shift=0.0, growth=None):
    return build(brief or brief0, v or v0, growth=growth, wacc_shift=wacc_shift)


def ps(r):
    return r["equity"]["per_share_gbp_pence"]


def pence(gbp_m):
    """GBP m today -> pence per share: GBP m / shares x 100 (desk.units.pounds_to_pence)."""
    return units.pounds_to_pence(gbp_m / shares)


base = run()
B = ps(base)
d = base["dcf"]
assert abs(B - engine_json["base"]["equity"]["per_share_gbp_pence"]) < 1e-9   # same as the engine's valuation.json
W0 = base["wacc"]["wacc"]
KE0 = base["wacc"]["cost_of_equity"]

# ================= Excel: recalculate in LibreOffice, compare with Python =================
excel = {"status": "LibreOffice not found: not checked"}
if shutil.which("soffice"):
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", tmp, str(HERE / "dcf.xlsx")],
                       check=True, capture_output=True)
        wb = load_workbook(Path(tmp) / "dcf.xlsx", data_only=True)
    s = wb["DCF"]
    labels = {s.cell(row=r, column=1).value: r for r in range(1, s.max_row + 1)}
    diffs = []
    row_key = {"Net sales": "sales", "EBIT": "ebit", "Tax": "tax", "NOPAT": "nopat", "D&A": "da", "Capex": "capex",
               "Increase in NWC": "delta_nwc", "Unlevered FCF": "ufcf", "Cash flow counted": "cash_flow_counted",
               "Discount period": "period", "Discount factor": "discount_factor", "PV of FCF": "pv"}
    for lab, key in row_key.items():
        for i, row in enumerate(d["rows"]):
            xl = s.cell(row=labels[lab], column=3 + i).value
            diffs.append((f"{lab} FY{row['year'] % 100}", xl, row[key]))
    summary_py = {"Cost of equity = Rf + beta x ERP": KE0, "After-tax cost of debt": base["wacc"]["after_tax_cost_of_debt"],
                  "WACC": W0, "Sum of PV of FCF": d["pv_forecast"],
                  "Terminal-year FCF (normalised, grows at g)": d["terminal_year"]["ufcf"],
                  "Terminal value (Gordon growth)": d["terminal_value"], "PV of terminal value": d["pv_terminal_value"],
                  "Enterprise value": d["enterprise_value"], "Minority interests deducted": base["equity"]["minorities_deducted"],
                  "Equity value at balance sheet date": base["equity"]["equity_value_balance_sheet_date"],
                  "Equity value at valuation date": base["equity"]["equity_value"],
                  "Value per share (GBP)": base["equity"]["per_share_gbp"], "Value per share (pence)": B,
                  "Upside vs share price": base["rating"]["upside"]}
    for lab, py in summary_py.items():
        diffs.append((lab, s.cell(row=labels[lab], column=2).value, py))
    rel = [(n, xl, py, abs(xl - py) / max(abs(py), 1e-12)) for n, xl, py in diffs]
    worst = max(rel, key=lambda t: t[3])
    tol = brief0["checks"]["excel_match_tolerance"]
    excel = {"cells_compared": len(rel), "worst_line": worst[0], "worst_relative_difference": worst[3],
             "tolerance": tol, "pass": worst[3] <= tol,
             "headline": {n: {"excel": xl, "python": py} for n, xl, py, _ in rel
                          if n in ("Enterprise value", "Value per share (pence)", "WACC", "Upside vs share price")}}
    assert excel["pass"]

# ================= sanity checks =================
peers = comps_out["peers"]
mults = [p["ev_ebitda_ltm"] for p in peers]
primary = comps_out["comps_range_pence_primary"]
checks = {
    "tv_share_of_ev": {"value": d["tv_share_of_ev"], "threshold": 0.75, "flag": d["tv_share_of_ev"] > 0.75},
    "implied_exit_ev_ebitda": {"value": d["implied_exit_ev_ebitda"], "peer_low": min(mults), "peer_high": max(mults),
                               "peer_median": comps_out["legs"]["ev_ebitda_ltm_x_calendarised_ebitda (PRIMARY)"]["median_multiple"],
                               "informa_optional": comps_out["optional_peers_shown_separately"][0]["ev_ebitda_ltm"],
                               "relx_market_today": comps_out["relx_own_multiples"]["market_ev_ebitda_actual_ltm"],
                               "within_peer_range": min(mults) <= d["implied_exit_ev_ebitda"] <= max(mults)},
    "wacc": {"value": W0, "range": [0.06, 0.09], "flag": not 0.06 <= W0 <= 0.09},
    "dcf_vs_comps": {"dcf_pence": B, "comps_median_pence": primary["implied_value"], "comps_low_pence": primary["low"],
                     "comps_high_pence": primary["high"], "gap": B / primary["implied_value"] - 1}}   # gap = DCF / comps - 1

# EV/EBITDA the DCF implies today (EV at 31 Dec 2025 / FY25 adjusted EBITDA), against the market's 14.6x
fy25_ebitda = comps_out["relx_metrics_mgbp"]["ebitda_fy25_adjusted"]
checks["dcf_ev_over_fy25_ebitda"] = d["enterprise_value"] / fy25_ebitda


# ================= flip points (one variable at a time) =================
def solve(fn, lo, hi, target, tol=1e-10):
    """Bisection: the parameter at which value per share = target (value assumed monotonic in the parameter)."""
    flo, fhi = fn(lo) - target, fn(hi) - target
    if flo * fhi > 0:
        return None
    for _ in range(300):
        mid = (lo + hi) / 2
        fm = fn(mid) - target
        if abs(hi - lo) < tol:
            break
        if fm * flo > 0:
            lo, flo = mid, fm
        else:
            hi = mid
    return (lo + hi) / 2


def with_v(**kw):
    v = copy.deepcopy(v0)
    v.update(kw)
    return v


def with_d(**kw):
    b = copy.deepcopy(brief0)
    b["dcf"].update(kw)
    return b


def erp_brief(x):
    b = copy.deepcopy(brief0)
    b["dcf"]["wacc"]["equity_risk_premium"] = x
    return b


# each runner returns the engine result for one changed input (everything else as in the base case)
RUNNERS = {
    "wacc_shift_added_to_wacc": lambda x: run(wacc_shift=x),
    "beta": lambda x: run(with_v(beta=x)),
    "base_margin": lambda x: run(with_v(base_margin=x)),
    # PM rule: pre-tax cost of debt = risk-free + credit spread, so the gilt moves both the cost of equity and of debt
    "risk_free_gilt": lambda x: run(with_v(risk_free=x, pre_tax_cost_of_debt=x + SPREAD)),
    "equity_risk_premium": lambda x: run(brief=erp_brief(x)),
    "terminal_growth": lambda x: run(brief=with_d(terminal_growth=x)),
    "growth_fy27_30": lambda x: run(brief=with_d(organic_growth=[brief0["dcf"]["organic_growth"][0]] + [x] * 4)),
}


def value_fn(label):
    return lambda x: ps(RUNNERS[label](x))


by_wacc_shift, by_beta, by_margin, by_rf, by_erp, by_tg, by_growth_fy27_30 = (value_fn(k) for k in RUNNERS)

flips = {}
for label, lo, hi, now in [("wacc_shift_added_to_wacc", -0.04, 0.02, 0.0), ("beta", 0.0, 1.5, v0["beta"]),
                           ("base_margin", 0.30, 0.70, v0["base_margin"]), ("risk_free_gilt", 0.0, 0.08, v0["risk_free"]),
                           ("equity_risk_premium", 0.0, 0.08, ERP), ("terminal_growth", 0.0, W0 - 0.0005, G0),
                           ("growth_fy27_30", 0.0, 0.25, brief0["dcf"]["organic_growth"][1])]:
    fn = value_fn(label)
    out = {"base_value": now}
    for name, tgt in (("sell_to_hold", hold_floor), ("hold_to_buy", buy_floor)):
        x = solve(fn, lo, hi, tgt)
        out[name] = {"parameter": x, "value_check_pence": fn(x) if x is not None else None,
                     "status": "found" if x is not None else f"not reached in [{lo}, {hi}]"}
        if x is not None:
            out[name]["wacc_level"] = RUNNERS[label](x)["wacc"]["wacc_used"]
    flips[label] = out
# WACC at which the DCF equals today's price (market-implied WACC)
x_mkt = solve(by_wacc_shift, -0.04, 0.02, price)
flips["market_implied_wacc"] = {"wacc": W0 + x_mkt, "value_check_pence": by_wacc_shift(x_mkt)}

# ================= value per 1 point of each key input (+1 and -1 percentage point, one at a time) =================
def dv(fn, now, step):
    return {"plus": fn(now + step) - B, "minus": fn(now - step) - B}


per_point = {
    "wacc (directly, 1 pt)": dv(by_wacc_shift, 0.0, 0.01),
    "risk_free_gilt (1 pt, moves Ke and Kd)": dv(by_rf, v0["risk_free"], 0.01),
    "equity_risk_premium (1 pt)": dv(by_erp, ERP, 0.01),
    "beta (0.10)": dv(by_beta, v0["beta"], 0.10),
    "operating_margin (1 pt)": dv(by_margin, v0["base_margin"], 0.01),
    "tax_rate (1 pt)": dv(lambda x: ps(run(with_v(tax_rate=x))), v0["tax_rate"], 0.01),
    "growth_fy27_30 (1 pt each year)": dv(by_growth_fy27_30, brief0["dcf"]["organic_growth"][1], 0.01),
    "growth_fy26_30 (1 pt each year)": dv(lambda x: ps(run(brief=with_d(organic_growth=[g + x for g in brief0["dcf"]["organic_growth"]]))), 0.0, 0.01),
    "terminal_growth (1 pt)": dv(by_tg, G0, 0.01),
    "capex_pct_sales (1 pt, forecast years; terminal set by RONIC)": dv(lambda x: ps(run(with_v(capex_pct_sales=x))), v0["capex_pct_sales"], 0.01),
    "da_pct_sales (1 pt)": dv(lambda x: ps(run(with_v(da_pct_sales=x))), v0["da_pct_sales"], 0.01),
    "nwc_pct_sales (1 pt)": dv(lambda x: ps(run(with_v(nwc_pct_sales=x))), v0["nwc_pct_sales"], 0.01),
    "terminal_ronic (1 pt)": dv(lambda x: ps(run(brief=with_d(terminal_ronic=x))), brief0["dcf"]["terminal_ronic"], 0.01),
    "debt_weight (1 pt, beta not relevered)": dv(lambda x: ps(run(with_v(debt_weight=x))), v0["debt_weight"], 0.01),
}

# ================= beta: PM card 0.974 vs run 0.9729 =================
card = run(with_v(beta=0.974))
beta_diff = {"run_beta": v0["beta"], "card_beta": 0.974,
             "run_value_pence": B, "card_value_pence": ps(card), "difference_pence": ps(card) - B,
             "run_wacc": W0, "card_wacc": card["wacc"]["wacc"], "card_cost_of_equity": card["wacc"]["cost_of_equity"],
             "card_rating": card["rating"]["rating"], "card_upside": card["rating"]["upside"],
             "note": "Only beta changes (debt weight, net debt and shares as in inputs.yaml)."}
# regression and other betas from the Risk agent, for context
beta_table = {}
for k, b in {"bottom_up_base": v0["beta"], "card_0.974": 0.974, "peer_median_raw": 0.75, "own_5y_weekly": 0.66,
             "own_2y_weekly": 0.67, "own_blume_5y_monthly": 0.58, "own_5y_monthly": 0.38}.items():
    r = run(with_v(beta=b))
    beta_table[k] = {"beta": b, "cost_of_equity": r["wacc"]["cost_of_equity"], "wacc": r["wacc"]["wacc"],
                     "value_pence": ps(r), "upside": r["rating"]["upside"], "rating": r["rating"]["rating"],
                     "tv_share": r["dcf"]["tv_share_of_ev"], "exit_ev_ebitda": r["dcf"]["implied_exit_ev_ebitda"]}

# ================= where the value per share comes from =================
f = dcf.roll_forward(1.0, KE0, v0["stub_years"])   # growth factor at Ke from 31 Dec 2025 to 13 Oct 2026
parts = {"pv_forecast_fcf_fy26_30": pence(d["pv_forecast"] * f), "pv_terminal_value": pence(d["pv_terminal_value"] * f),
         "less_net_debt_31dec2025": -pence(v0["net_debt"] * f), "less_minorities": -pence(v0["minority_interests"] * f),
         "plus_investments": pence(v0["associates"] * f),
         "less_dividends_since_31dec": -pence(v0["dividends_ex_before_valuation_mgbp"]),
         "less_buybacks_since_31dec": -pence(v0["buybacks_after_balance_sheet_mgbp"])}
assert abs(sum(parts.values()) - B) < 1e-6
decomp = {"pence": parts, "growth_factor_at_ke": f}

# ================= bear (engine) and sensitivity grid (engine's valuation.json) =================
bear = engine_json["bear"]
out = {"units": "value per share in pence (GBp) unless labelled; money in GBP m",
       "base_value_pence": B, "price_pence": price, "upside": base["rating"]["upside"], "rating": base["rating"]["rating"],
       "rating_thresholds_pence": {"hold_if_at_least": hold_floor, "buy_if_at_least": buy_floor},
       "excel_check": excel, "checks": checks, "flip_points": flips, "value_per_1pt_pence": per_point,
       "beta_card_vs_run": beta_diff, "beta_table": beta_table, "value_decomposition": decomp,
       "bear_engine": {"value_pence": bear["equity"]["per_share_gbp_pence"], "rating": bear["rating"]["rating"],
                       "upside": bear["rating"]["upside"], "wacc": bear["wacc"]["wacc_used"]},
       "terminal": {"tv_mgbp": d["terminal_value"], "pv_tv_mgbp": d["pv_terminal_value"],
                    "terminal_capex_pct_sales_from_ronic": d["terminal_year"]["capex_pct_sales"],
                    "reinvestment_rate": d["terminal_year"]["reinvestment_rate"], "terminal_fcf_mgbp": d["terminal_year"]["ufcf"]}}
(HERE / "sanity_checks.json").write_text(json.dumps(out, indent=2, default=float))
print(json.dumps(out, indent=1, default=float))
