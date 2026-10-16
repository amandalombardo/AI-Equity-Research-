"""Risk agent, RELX PLC (REL.L), run 2026-10-14: the valuation beta (PM rule) and its evidence.

THE ONLY SCRIPT THAT WRITES `beta` TO inputs.yaml (Codex audit RISK-01). It follows the brief's rule
(briefs/RELX.yaml dcf.wacc.beta, agreed 14 Oct 2026): bottom-up beta = Damodaran GLOBAL Information Services unlevered
beta corrected for cash, relevered (Hamada) at RELX's market D/E = 31 Dec 2025 net debt / market cap, and the brief's
22.8% tax rate. Regressions and peer betas (risk.json, from run_risk.py) are cross-checks only and never overwrite it.

Run from the repo root (after run_risk.py):
  PYTHONPATH=. python runs/RELX/2026-10-14/risk/beta_evidence.py <path to betaGlobal.xlsx>
The .xlsx is the PM's betaGlobal.xls converted with: soffice --headless --convert-to xlsx --outdir <scratch> <file>
"""
import hashlib
import json
import re
import sys
from pathlib import Path

import openpyxl
import yaml

from desk import dcf

HERE = Path(__file__).resolve().parent
RUN = HERE.parent
REPO = RUN.parents[2]
DAMODARAN_XLS = "docs/data/betaGlobal.xls"  # uploaded by the PM; same file as https://pages.stern.nyu.edu/~adamodar/pc/datasets/betaGlobal.xls
SHA256_XLS = "10c33f46df2a35fa7e4e8c4386c01f9da636fc2ab3e1e98307a52a947c9ab52b"
INDUSTRY, PM_VALUE, PM_FIRMS = "Information Services", 0.873735946265916, 86   # PM message, 14 Oct 2026

brief = yaml.safe_load((REPO / "briefs" / "RELX.yaml").read_text())
inp = yaml.safe_load((RUN / "inputs.yaml").read_text())
tax = brief["dcf"]["tax_rate"]                       # 22.8%, PM decision
erp = brief["dcf"]["wacc"]["equity_risk_premium"]    # 5.0%
rf = inp["risk_free"]["value"]                       # 5.45% gilt, 13 Oct 2026 (Data agent, from the brief)

if Path(DAMODARAN_XLS).exists() and hashlib.sha256(Path(DAMODARAN_XLS).read_bytes()).hexdigest() != SHA256_XLS:
    sys.exit("STOP: betaGlobal.xls has changed since the PM's upload (sha256 mismatch)")

# ---- 1. read the industry betas from the PM's file (do not trust a typed-in number)
wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
ws = wb["Industry Averages"]
file_date = ws["B1"].value
header, rows = None, {}
for r in ws.iter_rows(values_only=True):
    if r[0] == "Industry Name":
        header = [str(h).strip() if h else h for h in r]
    elif header and r[0]:
        rows[r[0]] = dict(zip(header, r))
row = rows.get(INDUSTRY)
if row is None:
    sys.exit(f"MISSING: {INDUSTRY} row not found in 'Industry Averages'. STOP.")
bu = float(row["Unlevered beta corrected for cash"])
if abs(bu - PM_VALUE) > 1e-12 or row["Number of firms"] != PM_FIRMS or str(file_date.date()) != "2026-01-05":
    sys.exit(f"STOP: file value {bu}, firms {row['Number of firms']}, date {file_date} do not match the PM's rule")
industry = {"industry": INDUSTRY, "region": "Global", "number_of_firms": row["Number of firms"],
            "file_date": str(file_date.date()), "levered_beta": row["Beta"], "industry_d_to_e": row["D/E Ratio"],
            "unlevered_beta": row["Unlevered beta"], "cash_to_firm_value": row["Cash/Firm value"],
            "unlevered_beta_corrected_for_cash": bu,
            "over_time": {h: row[h] for h in header if h and re.fullmatch(r"20\d\d|Average.*", h)},
            "source": f"Damodaran, Betas by Sector (Global), betaGlobal.xls dated 2026-01-05, sheet 'Industry Averages', "
                      f"row '{INDUSTRY}' (PM upload, {DAMODARAN_XLS}, sha256 {SHA256_XLS[:12]}...)"}
other_industries = {n: {"firms": rows[n]["Number of firms"], "unlevered_beta_corrected_for_cash": rows[n]["Unlevered beta corrected for cash"]}
                    for n in ("Publishing & Newspapers", "Business & Consumer Services")}

# ---- 2. RELX's market debt-to-equity, all in GBP m
shares_m = inp["shares_outstanding_m"]["value"]          # 1,743.754727m ex treasury, 25 Sep 2026 (6-K 0000929869-26-000071)
price_p = inp["share_price_gbp_pence"]["value"]          # 2,771p, REL.L close 13 Oct 2026
mcap = shares_m * price_p / 100                          # market cap in GBP m = shares x price in pounds (pence / 100)
nd = inp["net_debt"]["value"]                            # 7,104 = 7,201 - 97 leases, 31 Dec 2025 (FY25 results p.13)


def relever(unlevered: float, debt: float, equity: float, t: float) -> float:
    """Hamada: levered beta = unlevered beta x (1 + (1 - tax) x D/E). No desk function exists for this (FLAG)."""
    return unlevered * (1 + (1 - t) * debt / equity)


def ke(beta: float) -> float:
    """Cost of equity via desk.dcf.cost_of_equity (CAPM: risk-free + beta x equity risk premium)."""
    return dcf.cost_of_equity(rf, beta, erp)


base = relever(bu, nd, mcap, tax)
pm_card_mcap = 1747.0 * price_p / 100                    # PM card: 2,771p x 1,747.0m shares = GBP 48.4bn
cases = {
    "BASE: Information Services 0.8737, net debt ex leases 7,104 (31 Dec 2025)": (bu, nd, mcap),
    "PM card reconciliation: net debt incl. leases 7,201, this market cap": (bu, 7201.0, mcap),
    "PM card as printed: 7,201 and 1,747.0m shares": (bu, 7201.0, pm_card_mcap),
    "Sensitivity: net debt ex leases 30 Jun 2026 (8,614)": (bu, 8614.0, mcap),
    "Sensitivity: net debt RELX definition 30 Jun 2026 (8,733; PM option C)": (bu, 8733.0, mcap),
    "Sensitivity: unlevered beta without cash correction (0.8265)": (row["Unlevered beta"], nd, mcap),
    "Sensitivity: Damodaran 2020-24 average, cash-corrected (1.0729)": (row["Average (2020-24)"], nd, mcap),
    "Sensitivity: Publishing & Newspapers (0.7788)": (other_industries["Publishing & Newspapers"]["unlevered_beta_corrected_for_cash"], nd, mcap),
    "Sensitivity: Business & Consumer Services (0.9083)": (other_industries["Business & Consumer Services"]["unlevered_beta_corrected_for_cash"], nd, mcap),
}
relevered = {k: {"unlevered": u, "net_debt_mgbp": d, "market_cap_mgbp": e, "d_to_e": d / e,
                 "beta": relever(u, d, e, tax), "cost_of_equity": ke(relever(u, d, e, tax))} for k, (u, d, e) in cases.items()}

rj = json.loads((HERE / "risk.json").read_text())
own = rj["own_regressions_vs_ftse100"]
xchecks = {f"REL.L vs FTSE 100, {k.replace('_', ' ')}": {kk: own[k][kk] for kk in ("beta", "std_error", "ci95", "r_squared", "observations")}
           for k in own}
xchecks["Peer median (7 core peers), 5y monthly vs home index (raw, not relevered)"] = {"beta": rj["peer_median_beta_5y_monthly_core7"]}
inside = {k: v["ci95"][0] <= base <= v["ci95"][1] for k, v in xchecks.items() if "ci95" in v}
pm_ref = {"2y weekly": (0.67, 0.08), "5y weekly": (0.66, 0.12), "5y monthly": (0.38, 0.04)}   # docs/phase5/relx-beta-evidence.md
pm_match = {k: {"pm_beta": b, "pm_r2": r2, "agent_beta": own[k.replace(" ", "_")]["beta"], "agent_r2": own[k.replace(" ", "_")]["r_squared"],
                "match_to_2dp": round(own[k.replace(" ", "_")]["beta"], 2) == b and round(own[k.replace(" ", "_")]["r_squared"], 2) == r2}
            for k, (b, r2) in pm_ref.items()}

out = {
    "rule": "PM, briefs/RELX.yaml dcf.wacc.beta (agreed 14 Oct 2026): bottom-up Damodaran GLOBAL Information Services unlevered beta "
            "corrected for cash, relevered (Hamada) at RELX's market D/E (31 Dec 2025 net debt / market cap) and the 22.8% tax rate. "
            "Own regressions vs the FTSE 100 are the cross-check.",
    "industry": industry, "other_industries_reference": other_industries,
    "market_cap": {"shares_m": shares_m, "price_gbp_pence": price_p, "market_cap_gbp_m": mcap,
                   "sources": "inputs.yaml shares_outstanding_m (6-K 0000929869-26-000071 ex 99.6, 25 Sep 2026), "
                              "share_price_gbp_pence (Yahoo REL.L close 2026-10-13)"},
    "net_debt_mgbp_31_dec_2025_ex_leases": nd, "d_to_e": nd / mcap, "tax_rate": tax,
    "relevered_betas": relevered,
    "difference_vs_pm_card_0974": base - relevered["PM card as printed: 7,201 and 1,747.0m shares"]["beta"],
    "cross_checks": xchecks, "base_inside_each_95pct_interval": inside, "pm_evidence_regressions_match": pm_match,
    "recommended_beta": round(base, 4),
}
(HERE / "beta_evidence.json").write_text(json.dumps(out, indent=2, default=str))

# ---- 3. write ONLY `beta` into inputs.yaml, leaving every other entry's text untouched
src = (f"Risk agent, bottom-up: Damodaran GLOBAL Information Services unlevered beta corrected for cash {bu:.4f} ({PM_FIRMS} firms, "
       f"betaGlobal.xls dated 2026-01-05, sheet 'Industry Averages', PM upload) relevered (Hamada) at RELX's market D/E {nd / mcap:.4f} "
       f"= net debt ex leases GBP {nd:,.0f}m (31 Dec 2025, 7,201 less 97 leases, FY25 results 6-K 0000929869-26-000008 p.13) / "
       f"market cap GBP {mcap:,.1f}m ({shares_m:,.6f}m shares x {price_p:,.0f}p / 100) and tax {tax:.1%} (brief). PM decision 14 Oct 2026 "
       f"(briefs/RELX.yaml dcf.wacc.beta). The PM card's 0.974 used net debt incl. leases 7,201; at the brief's 7,104 the value is "
       f"{base:.4f}. Cross-checks in risk/beta_evidence.json and risk/risk.json (not used).")
text = (RUN / "inputs.yaml").read_text()
text = re.sub(r"(?ms)^beta:\n(?:[ ].*\n)*", "", text)          # drop any earlier beta block (idempotent re-runs)
text = text.rstrip("\n") + "\n" + yaml.safe_dump({"beta": {"value": float(round(base, 4)), "source": src}},
                                                   sort_keys=False, width=160)
check = yaml.safe_load(text)
assert {k: v for k, v in check.items() if k != "beta"} == {k: v for k, v in inp.items() if k != "beta"}, "other entries changed"
(RUN / "inputs.yaml").write_text(text)

print(f"{INDUSTRY} bu (cash-corrected) {bu}  firms {row['Number of firms']}  date {file_date.date()}")
print(f"mcap GBP m {mcap:,.1f}  D/E {nd / mcap:.5f}")
for k, v in relevered.items():
    print(f"{k:75s} D/E {v['d_to_e']:.4f} beta {v['beta']:.4f}  Ke {v['cost_of_equity']:.4%}")
print("inside 95% CI:", inside)
print("PM evidence match:", {k: v["match_to_2dp"] for k, v in pm_match.items()})
print("diff vs PM card", out["difference_vs_pm_card_0974"])
