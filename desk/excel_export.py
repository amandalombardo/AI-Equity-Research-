"""Write the base-case DCF to Excel with live formulas, so the PM can audit it and
compare it line by line with her own model. A 'Python check' column shows the value
the code calculated; if the Excel formula and the Python value disagree, one is wrong."""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from desk.units import price_currency, share_price_key

BOLD = Font(bold=True)
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")  # yellow = input


def write_dcf_workbook(path, brief: dict, v: dict, sources: dict, base: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Inputs"
    d = brief["dcf"]
    ccy = price_currency(brief["company"])
    brief_file = f"briefs/{brief['company'].get('ticker', 'DGE')}.yaml"
    inputs = [
        ("base_sales", "Base-year net sales, excluding businesses being sold (USD m)", v["base_sales"]),
        ("base_margin", "Base operating margin before exceptionals", v["base_margin"]),
        ("tax_rate", "Tax rate", v["tax_rate"]),
        ("da_pct", "D&A % of sales", v["da_pct_sales"]),
        ("capex_pct", "Capex % of sales", v["capex_pct_sales"]),
        ("tcapex_pct", "Terminal-year capex % of sales", base["dcf"]["terminal_year"]["capex_pct_sales"]),
        ("nwc_pct", "Net working capital % of sales", v["nwc_pct_sales"]),
        ("flow_through", "Savings flow-through", d["savings"]["flow_through"]),
        ("rf", "Risk-free rate (US 10y)", v["risk_free"]),
        ("erp", "Equity risk premium", d["wacc"]["equity_risk_premium"]),
        ("beta", "Beta", v["beta"]),
        ("kd", "Pre-tax cost of debt", v["pre_tax_cost_of_debt"]),
        ("wd", "Debt weight D/(D+E)", v["debt_weight"]),
        ("g", "Terminal growth", d["terminal_growth"]),
        ("stub", "Fraction of the first forecast year elapsed, used in discounting (0 = valued at the balance sheet date)",
         base["inputs"].stub_years),
        ("roll_years", "Years from the balance sheet date to valuation date (equity roll-forward)", (base["roll"] or {}).get("years", 0.0)),
        ("divs", "Dividends gone ex before valuation date (USD m)", (base["roll"] or {}).get("dividends", 0.0)),
        ("net_debt", "Net debt (USD m)", v["net_debt"]),
        ("leases", "Lease liabilities not in net debt (USD m)", v["lease_liabilities"]),
        ("minorities", "Minority interests at book (USD m)", v["minority_interests"]),
        ("mi_share", "Minorities' share of operating profit (proportionate method; blank = use book)",
         base["bridge"].get("minority_share_of_ev")),
        ("pension", "Pension deficit (USD m)", v["pension_deficit"]),
        ("associates", "Associates and equity stakes (USD m)", v["associates"]),
        ("disposals", "Disposal proceeds not in net debt (USD m)", v["disposal_proceeds"]),
        ("restructuring", "Restructuring cash still to pay (USD m)", v["restructuring_cash"]),
        ("contingent", "Contingent liabilities assumed paid, e.g. a lost tax case (USD m)", base["bridge"]["contingent_liabilities"]),
        ("shares", "Ordinary shares (m)", v["shares_outstanding_m"]),
    ]
    if ccy == "GBp":
        inputs += [("fx", "USD per GBP", v["fx_usd_per_gbp"]), ("price", "Share price (pence)", v["share_price_gbp_pence"])]
    else:
        inputs += [("price", "Share price (USD)", v["share_price_usd"])]
    ws.append(["Input", "Value", "Source"])
    for c in ws[1]:
        c.font = BOLD
    ref = {}
    tsrc = (f"PM rule: return on new investment {d['terminal_ronic']:.1%} (value driver formula, {brief_file})"
            if d.get("terminal_ronic") is not None else
            f"PM assumption ({brief_file})" if d.get("terminal_capex_pct_sales") is not None else
            "same as forecast capex % of sales")
    sources = {**sources, "tcapex_pct": tsrc}
    if base["roll"]:
        sources["stub"] = "0: valued at the balance sheet date, then rolled forward (brief: valuation_basis)"
    source_key = {"da_pct": "da_pct_sales", "capex_pct": "capex_pct_sales", "nwc_pct": "nwc_pct_sales", "rf": "risk_free",
                  "kd": "pre_tax_cost_of_debt", "wd": "debt_weight", "stub": "stub_years", "leases": "lease_liabilities",
                  "minorities": "minority_interests", "pension": "pension_deficit", "disposals": "disposal_proceeds", "restructuring": "restructuring_cash",
                  "mi_share": "minority_share_of_operating_profit",
                  "roll_years": "stub_years", "divs": "dividends_ex_before_valuation_musd",
                  "contingent": "contingent_liabilities",
                  "shares": "shares_outstanding_m", "fx": "fx_usd_per_gbp", "price": share_price_key(ccy)}
    for i, (key, label, val) in enumerate(inputs, start=2):
        src = sources.get(source_key.get(key, key)) or (
            "not in this brief's bridge (0)" if key == "contingent" and not val else f"PM assumption ({brief_file})")
        ws.append([label, val, src])
        ws.cell(row=i, column=2).fill = INPUT_FILL
        ref[key] = f"Inputs!$B${i}"
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["C"].width = 70

    s = wb.create_sheet("DCF")
    rows = base["dcf"]["rows"]
    n = len(rows)
    cols = [get_column_letter(3 + i) for i in range(n)]  # C.. for forecast years; B = base year
    base_fy = brief.get("base_year", {}).get("fiscal_year", rows[0]["year"] - 1)
    s.append(["USD m", f"FY{base_fy % 100} base"] + [f"FY{r['year'] % 100}" for r in rows] + ["", "Python check (final year)"])
    labels = ["Growth", "Net sales", "Savings kept", "EBIT", "Tax", "NOPAT", "D&A", "Capex", "Increase in NWC",
              "Unlevered FCF", "Cash flow counted", "Discount period", "Discount factor", "PV of FCF"]
    for lab in labels:
        s.append([lab])
    R = {lab: i + 2 for i, lab in enumerate(labels)}
    s[f"B{R['Net sales']}"] = f"={ref['base_sales']}"
    for i, c in enumerate(cols):
        prev = "B" if i == 0 else cols[i - 1]
        s[f"{c}{R['Growth']}"] = rows[i]["growth"]
        s[f"{c}{R['Growth']}"].fill = INPUT_FILL
        s[f"{c}{R['Net sales']}"] = f"={prev}{R['Net sales']}*(1+{c}{R['Growth']})"
        s[f"{c}{R['Savings kept']}"] = f"={d['savings']['cumulative_musd'][i]}*{ref['flow_through']}"
        s[f"{c}{R['EBIT']}"] = f"={c}{R['Net sales']}*{ref['base_margin']}+{c}{R['Savings kept']}"
        s[f"{c}{R['Tax']}"] = f"={c}{R['EBIT']}*{ref['tax_rate']}"
        s[f"{c}{R['NOPAT']}"] = f"={c}{R['EBIT']}-{c}{R['Tax']}"
        s[f"{c}{R['D&A']}"] = f"={c}{R['Net sales']}*{ref['da_pct']}"
        s[f"{c}{R['Capex']}"] = f"={c}{R['Net sales']}*{ref['capex_pct']}"
        s[f"{c}{R['Increase in NWC']}"] = f"=({c}{R['Net sales']}-{prev}{R['Net sales']})*{ref['nwc_pct']}"
        s[f"{c}{R['Unlevered FCF']}"] = (f"={c}{R['NOPAT']}+{c}{R['D&A']}-{c}{R['Capex']}-{c}{R['Increase in NWC']}")
        factor = f"(1-{ref['stub']})" if i == 0 else "1"
        s[f"{c}{R['Cash flow counted']}"] = f"={c}{R['Unlevered FCF']}*{factor}"
        if d["mid_year_convention"]:
            start = "0" if i == 0 else f"({i}-{ref['stub']})"
            s[f"{c}{R['Discount period']}"] = f"=({start}+({i + 1}-{ref['stub']}))/2"
        else:
            s[f"{c}{R['Discount period']}"] = f"={i + 1}-{ref['stub']}"
        s[f"{c}{R['Discount factor']}"] = f"=1/(1+WACC)^{c}{R['Discount period']}"
        s[f"{c}{R['PV of FCF']}"] = f"={c}{R['Cash flow counted']}*{c}{R['Discount factor']}"
    chk = get_column_letter(3 + n + 1)
    last = rows[-1]
    for lab, key in [("Net sales", "sales"), ("EBIT", "ebit"), ("Unlevered FCF", "ufcf"), ("PV of FCF", "pv")]:
        s[f"{chk}{R[lab]}"] = last[key]

    lc = cols[-1]
    r0 = len(labels) + 3
    w = base["wacc"]
    eq = base["equity"]
    sales_l, nopat_l = f"{lc}{R['Net sales']}", f"{lc}{R['NOPAT']}"
    term = base["dcf"]["terminal_year"]
    spec = [  # (key, label, formula with {key} placeholders for other rows, Python value)
        ("ke", "Cost of equity = Rf + beta x ERP", f"={ref['rf']}+{ref['beta']}*{ref['erp']}", w["cost_of_equity"]),
        ("kd", "After-tax cost of debt", f"={ref['kd']}*(1-{ref['tax_rate']})", w["after_tax_cost_of_debt"]),
        ("wacc", "WACC", f"=(1-{ref['wd']})*{{ke}}+{ref['wd']}*{{kd}}", w["wacc_used"]),
        ("pv", "Sum of PV of FCF", f"=SUM(C{R['PV of FCF']}:{lc}{R['PV of FCF']})", base["dcf"]["pv_forecast"]),
        # normalised steady-state year after the forecast: everything grows at g; NWC rises only by what g needs
        ("tfcf", "Terminal-year FCF (normalised, grows at g)",
         (f"={nopat_l}*(1+{ref['g']})+{sales_l}*(1+{ref['g']})*{ref['da_pct']}"
          f"-{sales_l}*(1+{ref['g']})*{ref['tcapex_pct']}-{sales_l}*{ref['g']}*{ref['nwc_pct']}"), term["ufcf"]),
        ("tv", "Terminal value (Gordon growth)", f"={{tfcf}}/(WACC-{ref['g']})", base["dcf"]["terminal_value"]),
        ("pvtv", "PV of terminal value", f"={{tv}}/(1+WACC)^({n}-{ref['stub']}{'-0.5' if d['mid_year_convention'] else ''})", base["dcf"]["pv_terminal_value"]),
        ("ev", "Enterprise value", "={pv}+{pvtv}", base["dcf"]["enterprise_value"]),
        # proportionate minorities: outside shareholders' share of operating profit x EV; else book value
        ("mi", "Minority interests deducted", f"=IF(ISBLANK({ref['mi_share']}),{ref['minorities']},{ref['mi_share']}*{{ev}})",
         eq["minorities_deducted"]),
        ("eqbs", "Equity value at balance sheet date", (f"={{ev}}-{ref['net_debt']}-{ref['leases']}-{{mi}}"
                                f"-{ref['pension']}-{ref['restructuring']}-{ref['contingent']}+{ref['associates']}+{ref['disposals']}"),
         eq["equity_value_balance_sheet_date"]),
        # roll forward to the valuation date at the cost of equity, less dividends gone ex in between
        ("eq", "Equity value at valuation date", f"={{eqbs}}*(1+{{ke}})^{ref['roll_years']}-{ref['divs']}", eq["equity_value"]),
        ("usd", "Value per share (USD)", f"={{eq}}/{ref['shares']}", eq["per_share_usd"]),
    ]
    if ccy == "GBp":
        spec += [("gbp", "Value per share (pence)", f"={{usd}}/{ref['fx']}*100", eq["per_share_gbp_pence"]),
                 ("up", "Upside vs share price", f"={{gbp}}/{ref['price']}-1", base["rating"]["upside"])]
    else:
        spec += [("up", "Upside vs share price", f"={{usd}}/{ref['price']}-1", base["rating"]["upside"])]
    cell = {key: f"B{r0 + i}" for i, (key, *_rest) in enumerate(spec)}
    summary = [(lab, formula.format(**cell), py) for _key, lab, formula, py in spec]
    s[f"A{r0 - 1}"] = "Valuation"
    s[f"A{r0 - 1}"].font = BOLD
    s[f"C{r0 - 1}"] = "Python check"
    s[f"C{r0 - 1}"].font = BOLD
    for i, (lab, formula, py) in enumerate(summary):
        s[f"A{r0 + i}"] = lab
        s[f"B{r0 + i}"] = formula
        s[f"C{r0 + i}"] = py
    from openpyxl.workbook.defined_name import DefinedName
    wb.defined_names["WACC"] = DefinedName("WACC", attr_text=f"DCF!${cell['wacc'][0]}${cell['wacc'][1:]}")
    s.column_dimensions["A"].width = 36
    for row in s.iter_rows(min_row=2):
        for c in row[1:]:
            if isinstance(c.value, (int, float)) or (isinstance(c.value, str) and c.value.startswith("=")):
                c.number_format = "#,##0.0000" if c.row in (R["Growth"], R["Discount factor"], R["Discount period"]) else "#,##0.0"
    for i in (0, 1, 2, len(spec) - 1):
        s[f"B{r0 + i}"].number_format = s[f"C{r0 + i}"].number_format = "0.00%"
    wb.save(path)
